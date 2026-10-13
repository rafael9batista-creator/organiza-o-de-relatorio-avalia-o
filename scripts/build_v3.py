# -*- coding: utf-8 -*-
"""v3: acrescenta a aba SINTESE (uma célula sintética por conjunto) à planilha v2.
Uso: python3 -I build_v3.py <v2.xlsx> <saida_v3.xlsx>"""
import sys
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
sys.path.insert(0, __file__.rsplit('/', 1)[0])
from sintese_v3 import SINTESE, K, B, R, G, Y, N, P

SRC, OUT = sys.argv[1], sys.argv[2]
thin = Side(style='thin', color='BFBFBF'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP = Alignment(wrap_text=True, vertical='top')

def rich(parts):
    blocks = []
    for item in parts:
        if item[0] == 'H':
            f = InlineFont(rFont='Arial', sz=9, b=True, color=Color(rgb='FF' + K)); blocks.append(TextBlock(f, item[1]))
        else:
            col, txt = item
            blocks.append(TextBlock(InlineFont(rFont='Arial', sz=9, color=Color(rgb='FF' + col)), txt))
    return CellRichText(blocks)

wb = openpyxl.load_workbook(SRC, rich_text=True)
if 'SINTESE' in wb.sheetnames:
    del wb['SINTESE']
ws = wb.create_sheet('SINTESE', 1)
ws['A1'] = 'SÍNTESE por conjunto (v3 – 13/10/2026): descrição do conjunto (crítica + placa), componentes (função, marca, modelo, capacidade, série), notas fiscais e cotação; divergências na coluna Observações; valores da crítica, das NFs e da cotação para comparação.'
ws['A1'].font = Font(name='Arial', size=10, bold=True, color='1F3864')
ws['A2'] = 'Cores: preto = crítica de orçamento | azul = placa/campo/ANP | vermelho = NF (RT0538) | verde = SAP | amarelo = projeto | azul escuro = cotação OPE154 | roxo = avaliador (divergências, correções, proposições). Sem texto tachado: as redundâncias foram eliminadas na redação. Valores em branco = não atribuível ao subitem.'
ws['A2'].font = Font(name='Arial', size=9, color=P)
heads = ['Linha', 'Nível', 'Conjunto', 'Síntese (descrição, componentes, NF e cotação)', 'Observações / divergências', 'Valor crítica (R$)', 'Valor NF (R$)', 'Valor cotação (R$)', 'NF − crítica (R$)', 'Cotação − crítica (R$)', 'Fotos / fontes de campo']
for i, h in enumerate(heads, 1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = Font(name='Arial', size=9, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BORDER
for i, w in enumerate([8, 12, 32, 150, 60, 15, 15, 15, 14, 14, 30], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[4].height = 30
r = 5
for linha, nivel, nome, sint, obs, vc, vnf, vcot, fotos in SINTESE:
    ws.cell(row=r, column=1, value=linha); ws.cell(row=r, column=2, value=nivel); ws.cell(row=r, column=3, value=nome)
    ws.cell(row=r, column=4, value=rich(sint)); ws.cell(row=r, column=5, value=rich(obs))
    if vc is not None: ws.cell(row=r, column=6, value=vc)
    if vnf is not None: ws.cell(row=r, column=7, value=vnf).font = Font(name='Arial', size=9, color=R)
    if vcot is not None: ws.cell(row=r, column=8, value=vcot).font = Font(name='Arial', size=9, color=N)
    if vc is not None and vnf is not None: ws.cell(row=r, column=9, value=f'=G{r}-F{r}')
    if vc is not None and vcot is not None: ws.cell(row=r, column=10, value=f'=H{r}-F{r}')
    ws.cell(row=r, column=11, value=fotos)
    for ci in range(1, 12):
        c = ws.cell(row=r, column=ci); c.alignment = TOP; c.border = BORDER
        if ci in (1, 2, 3, 11): c.font = Font(name='Arial', size=9, bold=(ci == 3))
        if ci in (6, 9, 10): c.font = Font(name='Arial', size=9)
        if 6 <= ci <= 10: c.number_format = '#,##0.00;[Red]-#,##0.00'
    ntxt = sum(len(t) for _, t in sint)
    ws.row_dimensions[r].height = min(409, max(60, 12 * (ntxt // 140 + sint.__len__())))
    r += 1
last = r - 1
ws.cell(row=r + 1, column=3, value='Totais (linhas CONJUNTO; grupos 14 e 15-17 pelos subitens)').font = Font(name='Arial', size=9, bold=True)
ws.cell(row=r + 1, column=6, value=f'=SUMIFS(F5:F{last},B5:B{last},"CONJUNTO")')
ws.cell(row=r + 1, column=7, value=f'=SUMIFS(G5:G{last},B5:B{last},"CONJUNTO")+SUMIFS(G5:G{last},B5:B{last},"GRUPO")-G{[i for i,s in enumerate(SINTESE,5) if s[0]=="14d"][0]}')
ws.cell(row=r + 1, column=8, value=f'=SUMIFS(H5:H{last},B5:B{last},"CONJUNTO")+H{[i for i,s in enumerate(SINTESE,5) if s[0]=="14"][0]}')
ws.cell(row=r + 2, column=3, value='Conferência: crítica deve dar R$ 16.576.852,44; NF (soma das 23 NFs) R$ 16.782.799,46 – a coluna G soma as NFs sem dupla contagem (grupo 14 inclui a 1551 do 14d; grupo 15-17 cobre 15/16/17).').font = Font(name='Arial', size=9, color=P)
for ci in (6, 7, 8):
    ws.cell(row=r + 1, column=ci).number_format = '#,##0.00'; ws.cell(row=r + 1, column=ci).font = Font(name='Arial', size=9, bold=True)
ws.freeze_panes = 'D5'; ws.auto_filter.ref = f'A4:K{last}'

# LEIA-ME
wsl = wb['LEIA-ME']
wsl['A1'].value = 'Caracterização consolidada de M&E – Ultracargo Palmeirante-TO (base 07/10/2026; v2 12/10 com cotações OPE154; v3 13/10 com aba SINTESE)'
n = wsl.max_row + 2
wsl.cell(row=n, column=1, value='v3 – 13/10/2026 – aba SINTESE: uma célula sintética por conjunto, redigida de novo (sem repetições e sem tachado), na ordem descrição do conjunto → componentes → NF e cotação; coluna própria de Observações/divergências; colunas de valor da crítica, das NFs e da cotação com as diferenças. A síntese privilegia a crítica de orçamento e as leituras de campo; projeto, SAP, NF e cotação entram só no que acrescentam. A aba CONDENSADO (v2, com tudo e tachado) foi mantida para rastreabilidade.').font = Font(name='Arial', size=10, color=P, bold=True)
wsl.cell(row=n, column=1).alignment = Alignment(wrap_text=True, vertical='top')
wb.save(OUT)
print('salvo', OUT, '| linhas SINTESE:', len(SINTESE))
for linha, nivel, nome, sint, obs, *_ in SINTESE:
    print(f'{linha:>10} {sum(len(t) for _, t in sint):5d} chars')
